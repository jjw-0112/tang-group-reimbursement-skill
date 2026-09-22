"""Behavioral checks for extraction, decisions, and versioned reruns."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from PIL import Image
from pypdf import PdfReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from engine import _item_lines, _ofd_text, scan  # noqa: E402
from workflow import _item_thresholds, _write_workbook, build, evaluate  # noqa: E402


BUYER = "西安交通大学"
TAX_ID = "12100000435230200R"
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))


def invoice_pdf(number: str, amount: str, product: str = "运费", buyer: str = BUYER,
                issue_date: str = "2026年06月01日", seller: str = "某供应商",
                seller_tax_id: str = "913000000000000001") -> bytes:
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf)
    pdf.setFont("STSong-Light", 11)
    lines = [
        "电子发票（普通发票） 发票号码：" + number,
        "开票日期：" + issue_date,
        "购 名称：" + buyer + "       销 名称：" + seller,
        "统一社会信用代码/纳税人识别号：" + TAX_ID + "    统一社会信用代码/纳税人识别号：" + seller_tax_id,
        "项目名称 规格型号 数量 单价 金额 税率 税额",
        f"*服务*{product} 1 {amount} {amount} 0% 0.00",
        "合计",
        "价税合计（小写）¥" + amount,
    ]
    for i, line in enumerate(lines):
        pdf.drawString(35, 790 - i * 24, line)
    pdf.save()
    return buf.getvalue()


class ReimbursementTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.input = self.root / "原件"
        self.output = self.root / "输出"
        self.input.mkdir()
        self.profile = {"name": "测试学生", "student_id": "123", "phone": "13800000000", "cutoff_date": "2026-06-15"}

    def draft(self) -> dict:
        result = scan(self.input, self.output)
        result["profile"] = self.profile
        return result

    def test_item_thresholds_apply_to_each_product(self) -> None:
        text = """项目名称 规格型号 数量 单价 金额 税率 税额
*材料*零件甲 2 150.00 300.00 0% 0.00
*材料*零件乙 1 300.00 300.00 0% 0.00
*材料*传感器 2 250.00 500.00 0% 0.00
*材料*探头 1 1000.00 1000.00 0% 0.00
合计"""
        items, verified = _item_lines(text, Decimal("2100.00"))
        self.assertTrue(verified)
        self.assertEqual(len(items), 4)
        record = {"items": items}
        inventory, high_value = _item_thresholds(record)
        self.assertEqual(len(inventory), 2)
        self.assertTrue(any("传感器" in item for item in inventory))
        self.assertFalse(any("零件甲" in item or "零件乙" in item for item in inventory))
        self.assertEqual(len(high_value), 1)
        self.assertIn("探头", high_value[0])

    def test_buyer_history_and_dates(self) -> None:
        (self.input / "1-材料费20.pdf").write_bytes(invoice_pdf("12345678901234567890", "20.00", "零件", "其他单位"))
        draft = self.draft()
        record = evaluate(draft)[0][0]
        self.assertEqual(record["status"], "ineligible")
        self.assertIn("抬头不符", "；".join(record["issues"]))
        draft["records"][0]["manual"]["buyer_name"] = BUYER
        draft["history_invoice_numbers"] = ["12345678901234567890"]
        record = evaluate(draft)[0][0]
        self.assertEqual(record["status"], "pending")
        self.assertIn("疑似重复", "；".join(record["issues"]))
        draft["history_invoice_numbers"] = []
        draft["records"][0]["manual"]["issue_date"] = "2025-12-31"
        record = evaluate(draft)[0][0]
        self.assertEqual(record["status"], "ready")
        self.assertIn("跨年度", "；".join(record["warnings"]))
        draft["records"][0]["manual"]["issue_date"] = "2026-06-16"
        self.assertEqual(evaluate(draft)[0][0]["status"], "pending")

    def test_public_payment_conflict_and_private_high_value(self) -> None:
        (self.input / "1-材料费20.pdf").write_bytes(invoice_pdf("12345678901234567890", "20.00", "零件"))
        (self.input / "1-付款凭证.txt").write_text("支付金额：99.00 元", encoding="utf-8")
        draft = self.draft()
        draft["public_invoice_numbers"] = ["12345678901234567890"]
        record = evaluate(draft)[0][0]
        self.assertEqual(record["status"], "pending")
        self.assertIn("付款凭证金额", "；".join(record["issues"]))
        draft["records"][0]["manual"]["confirm_payment_amount"] = True
        self.assertEqual(evaluate(draft)[0][0]["status"], "ready")
        draft["public_invoice_numbers"] = []
        draft["records"][0]["manual"]["claim_amount"] = "1000.00"
        self.assertEqual(evaluate(draft)[0][0]["status"], "pending")
        (self.input / "1-材料费20.pdf").write_bytes(invoice_pdf("12345678901234567890", "1000.00", "零件"))
        (self.input / "1-材料费20.pdf").rename(self.input / "2-材料费1000.pdf")
        draft = self.draft()
        record = evaluate(draft)[0][0]
        self.assertEqual(record["status"], "missing")
        self.assertFalse(record["public_transfer"])
        self.assertIn("高值耗材登记材料", "；".join(record["missing_docs"]))
        self.assertFalse(any("支付记录" in issue for issue in record["missing_docs"]))
        self.assertNotIn("不可报销", "；".join(record["issues"]))
        package = build(draft)
        self.assertEqual(package["formal_count"], 1)
        ws = load_workbook(Path(package["directory"]) / "测试学生-报销.xlsx").active
        self.assertEqual(ws["C4"].value, 1000)

    def test_invoice_total_over_1000_requires_payment_not_high_value_when_units_are_lower(self) -> None:
        (self.input / "大额材料.pdf").write_bytes(invoice_pdf("63345678901234567890", "1044.54", "零件"))
        draft = self.draft()
        record = draft["records"][0]
        record["items"] = [
            {"name": "金属零件甲", "quantity": "4", "unit_gross": "174.60", "line_gross": "698.38"},
            {"name": "金属零件乙", "quantity": "4", "unit_gross": "86.54", "line_gross": "346.16"},
        ]
        record["items_verified"] = True
        record["manual"]["category"] = "材料费"
        evaluated = evaluate(draft)[0][0]
        self.assertEqual(evaluated["status"], "missing")
        self.assertIn("入库单：金属零件甲", "；".join(evaluated["missing_docs"]))
        self.assertIn("支付记录（票面总价超 1000 元）", evaluated["missing_docs"])
        self.assertFalse(any("高值耗材" in issue for issue in evaluated["missing_docs"]))
        record["manual"]["claim_amount"] = "999.00"
        self.assertIn("支付记录（票面总价超 1000 元）", evaluate(draft)[0][0]["missing_docs"])

    def test_same_name_contiguous_and_same_seller_date_group_marked(self) -> None:
        sources = [
            ("1-screw.pdf", "16345678901234567890", "螺丝", "2026年06月05日", "甲供应商有限公司", "913000000000000001"),
            ("2-nut.pdf", "17345678901234567890", "螺母", "2026年06月05日", "甲供应商有限公司", "913000000000000001"),
            ("3-screw.pdf", "18345678901234567890", "螺丝", "2026年06月05日", "甲供应商有限公司", "913000000000000001"),
            ("4-connector.pdf", "19345678901234567890", "连接器", "2026年06月05日", "乙供应商有限公司", "913000000000000002"),
            ("5-screw.pdf", "20345678901234567890", "螺丝", "2026年06月06日", "甲供应商有限公司", "913000000000000001"),
        ]
        for ref, number, product, issued, seller, seller_id in sources:
            (self.input / ref).write_bytes(invoice_pdf(number, "10.00", product, issue_date=issued,
                                                       seller=seller, seller_tax_id=seller_id))
        draft = self.draft()
        for record in draft["records"]:
            record["manual"].update(category="材料费", description=next(p for ref, _, p, _, _, _ in sources if ref == record["ref"]))
            # Old review.json files lack these fields; grouping must still use the invoice text.
            record.pop("seller_name", None)
            record.pop("seller_tax_id", None)
        package = build(draft)
        state = json.loads((self.output / "state.json").read_text(encoding="utf-8"))
        formal = sorted((r for r in state["records"] if r["status"] in {"ready", "missing", "pending"}),
                        key=lambda r: r["number"])
        names = [r["description"] for r in formal]
        screw_positions = [i for i, name in enumerate(names) if name.startswith("螺丝")]
        self.assertEqual(screw_positions, list(range(screw_positions[0], screw_positions[0] + 3)))
        self.assertEqual([names[i] for i in screw_positions], ["螺丝1", "螺丝2", "螺丝3"])
        by_ref = {r["ref"]: r for r in formal}
        grouped = [by_ref[ref]["same_seller_date_group"] for ref in ("1-screw.pdf", "2-nut.pdf", "3-screw.pdf")]
        self.assertEqual(grouped, ["G1", "G1", "G1"])
        self.assertFalse(by_ref["4-connector.pdf"].get("same_seller_date_group"))
        self.assertFalse(by_ref["5-screw.pdf"].get("same_seller_date_group"))
        ws = load_workbook(Path(package["directory"]) / "测试学生-报销.xlsx").active
        remarks = {str(row[1].value): str(row[3].value) for row in ws if isinstance(row[0].value, int)}
        self.assertIn("同开票方同日组 G1（3张，甲供应商有限公司，2026-06-05）", remarks["螺母"])
        self.assertNotIn("同开票方同日组", remarks["螺丝3"])

    def test_missing_rerun_renumber_and_confirmed_replacement(self) -> None:
        first = self.input / "1-快递费50.pdf"
        first.write_bytes(invoice_pdf("12345678901234567890", "50.00"))
        original_hash = hashlib.sha256(first.read_bytes()).hexdigest()
        draft = self.draft()
        self.assertEqual(len(draft["records"]), 1)
        self.assertEqual(draft["records"][0]["amount"], "50.00")
        self.assertEqual(evaluate(draft)[0][0]["status"], "missing")
        package1 = build(draft)
        self.assertEqual(package1["version"], 1)
        draft["records"][0]["manual"]["notes"] = "质检修订"
        revised = build(draft)
        self.assertEqual(revised["version"], 1)
        self.assertTrue(revised["revised_same_run"])
        self.assertEqual(len(list((self.output / ".drafts" / draft["run_id"]).glob("attempt-*"))), 1)
        with ZipFile(package1["zip"]) as archive:
            names = archive.namelist()
            self.assertTrue(any("1-快递费-50元.pdf" in name for name in names))
            self.assertTrue(any("补件与复核提醒.pdf" in name for name in names))
        wb = load_workbook(Path(package1["directory"]) / "测试学生-报销.xlsx")
        ws = wb.active
        self.assertEqual(ws["B1"].value, "测试学生")
        self.assertTrue(any("缺件" in str(row[3].value) for row in ws.iter_rows(min_row=2, max_col=4) if row[3].value))
        self.assertTrue(any(cell.data_type == "f" for row in ws for cell in row))

        (self.input / "1-运单.txt").write_text("运单：本票运费 50 元", encoding="utf-8")
        draft = self.draft()
        self.assertEqual(evaluate(draft)[0][0]["status"], "ready")
        package2 = build(draft)
        self.assertEqual(package2["version"], 2)
        self.assertTrue(Path(package1["zip"]).is_file())

        (self.input / "0-材料费20.pdf").write_bytes(invoice_pdf("22345678901234567890", "20.00", "零件"))
        draft = self.draft()
        package3 = build(draft)
        with (Path(package3["directory"]) / "编号变更.csv").open(encoding="utf-8-sig", newline="") as handle:
            changes = list(csv.DictReader(handle))
        self.assertTrue(any(row["变更类型"] == "重新编号" and row["旧编号"] == "1" and row["新编号"] == "2" for row in changes))

        (self.input / "2-快递费52.pdf").write_bytes(invoice_pdf("12345678901234567890", "52.00"))
        draft = self.draft()
        self.assertEqual(len(draft["replacement_candidates"]), 1)
        new_id = draft["replacement_candidates"][0]["new_id"]
        old_id = draft["replacement_candidates"][0]["old_id"]
        self.assertTrue(all(r["status"] == "pending" for r in evaluate(draft)[0] if r["invoice_number_final"] == "12345678901234567890"))
        next(r for r in draft["records"] if r["id"] == new_id)["manual"] = {
            "replacement_for": old_id, "attachment_refs": ["1-运单.txt"]
        }
        package4 = build(draft)
        state = json.loads((self.output / "state.json").read_text(encoding="utf-8"))
        self.assertEqual(next(r for r in state["records"] if r["id"] == old_id)["status"], "replaced")
        self.assertEqual(next(r for r in state["records"] if r["id"] == new_id)["status"], "ready")
        formal = [r for r in state["records"] if r["status"] in {"ready", "missing", "pending"}]
        self.assertEqual(sorted(r["number"] for r in formal), list(range(1, len(formal) + 1)))
        with (Path(package4["directory"]) / "文件对照表.csv").open(encoding="utf-8-sig", newline="") as handle:
            mappings = list(csv.DictReader(handle))
        with ZipFile(package4["zip"]) as archive:
            names = set(archive.namelist())
            self.assertTrue(all("v4/" + row["输出文件"] in names for row in mappings if row["输出文件"]))
        with (Path(package4["directory"]) / "编号变更.csv").open(encoding="utf-8-sig", newline="") as handle:
            changes = list(csv.DictReader(handle))
        self.assertTrue(any(row["变更类型"] == "更正发票替换" and row["旧编号"] == "2" and row["新编号"] == "2" for row in changes))
        self.assertEqual(hashlib.sha256(first.read_bytes()).hexdigest(), original_hash)

    def test_zip_and_ofd(self) -> None:
        with ZipFile(self.input / "一批材料.zip", "w") as archive:
            archive.writestr("发票/1-材料费20.pdf", invoice_pdf("32345678901234567890", "20.00", "零件"))
        draft = self.draft()
        self.assertEqual(draft["records"][0]["amount"], "20.00")
        self.assertIn("::", draft["records"][0]["ref"])
        package = build(draft)
        self.assertEqual(package["formal_count"], 1)
        with ZipFile(package["zip"]) as archive:
            self.assertTrue(any(name.endswith(".pdf") and "/发票/" in name for name in archive.namelist()))
        ofd = io.BytesIO()
        with ZipFile(ofd, "w") as archive:
            archive.writestr("Doc_0/Pages/Page_0/Content.xml", "<Page><TextObject><TextCode>铁路电子客票 票价：289.00 发票号码：42345678901234567890 开票日期：2026年06月01日</TextCode></TextObject></Page>")
        text, error = _ofd_text(ofd.getvalue())
        self.assertIsNone(error)
        self.assertIn("289.00", text)

    def test_unreadable_image_remains_pending(self) -> None:
        Image.new("RGB", (220, 100), "white").save(self.input / "scan.png")
        draft = self.draft()
        self.assertEqual(len(draft["records"]), 1)
        self.assertEqual(len(draft["attachments"]), 0)
        self.assertEqual(evaluate(draft)[0][0]["status"], "pending")
        package = build(draft)
        self.assertEqual(package["formal_count"], 1)
        self.assertTrue(any("金额待核" in p.name for p in (Path(package["directory"]) / "发票").iterdir()))
        ws = load_workbook(Path(package["directory"]) / "测试学生-报销.xlsx").active
        self.assertIsNone(ws["C4"].value)
        self.assertIn("待核实", str(ws["D4"].value))
        self.assertFalse((Path(package["directory"]) / "待核实").exists())

    def test_daily_item_has_own_section_and_remains_provisional(self) -> None:
        (self.input / "拖把.pdf").write_bytes(invoice_pdf("72345678901234567890", "56.90", "平板拖把"))
        draft = self.draft()
        self.assertEqual(draft["records"][0]["category_guess"], "日用及两用物品")
        record = evaluate(draft)[0][0]
        self.assertEqual(record["status"], "missing")
        self.assertIn("科研用途待负责人确认", record["missing_docs"])
        package = build(draft)
        ws = load_workbook(Path(package["directory"]) / "测试学生-报销.xlsx").active
        self.assertEqual(ws["A2"].value, "日用及两用物品")
        self.assertIn("待确认：科研用途", str(ws["D4"].value))
        self.assertEqual(ws["C4"].value, 56.9)
        draft["records"][0]["manual"]["research_use_confirmed"] = True
        self.assertEqual(evaluate(draft)[0][0]["status"], "ready")

    def test_workbook_lists_linked_materials_and_uses_requested_fonts(self) -> None:
        (self.input / "网约车.pdf").write_bytes(invoice_pdf("82345678901234567890", "50.00", "客运服务"))
        draft = self.draft()
        record = draft["records"][0]
        record["manual"] = {
            "category": "交通费", "description": "网约车A", "notes": "核对PDF",
            "attachment_refs": ["行程.pdf", "付款.png", "论文录用通知.pdf"],
        }
        draft["attachments"] = [
            {"ref": "行程.pdf", "ext": ".pdf", "attachment_type": "itinerary", "text_excerpt": "", "scan_errors": []},
            {"ref": "付款.png", "ext": ".png", "attachment_type": "payment", "text_excerpt": "", "scan_errors": []},
            {"ref": "论文录用通知.pdf", "ext": ".pdf", "attachment_type": "acceptance", "text_excerpt": "", "scan_errors": []},
            {"ref": "未关联通知.pdf", "ext": ".pdf", "attachment_type": "fee_notice", "text_excerpt": "", "scan_errors": []},
        ]
        evaluated = evaluate(draft)[0]
        self.assertEqual(evaluated[0]["status"], "ready")
        evaluated[0]["number"] = 1
        path = self.root / "format.xlsx"
        _write_workbook(path, self.profile, evaluated)
        ws = load_workbook(path, rich_text=True).active
        remark = str(ws["D4"].value)
        self.assertIn("附件：", remark)
        for label in ("行程单", "支付记录截图", "论文录用通知"):
            self.assertIn(label, remark)
        self.assertIn("核对PDF", remark)
        self.assertNotIn("缴费通知", remark)
        self.assertEqual(ws["A2"].fill.fgColor.rgb, "FFE2EFDA")
        self.assertTrue(ws["A2"].font.bold)
        self.assertEqual(ws["A2"].font.color.rgb, "FF000000")
        self.assertEqual(ws["C4"].font.name, "Times New Roman")
        self.assertEqual(ws["C4"].font.color.rgb, "FF000000")
        self.assertIsInstance(ws["B4"].value, CellRichText)
        runs = [part for part in ws["B4"].value if isinstance(part, TextBlock)]
        self.assertEqual({part.font.rFont for part in runs}, {"黑体", "Times New Roman"})
        self.assertTrue(all(part.font.color.rgb == "FF000000" for part in runs))

    def test_chinese_words_separated_by_space_do_not_create_invalid_rich_run(self) -> None:
        (self.input / "箱子.pdf").write_bytes(invoice_pdf("92345678901234567890", "49.80", "收纳箱"))
        draft = self.draft()
        draft["records"][0]["manual"]["description"] = "收纳箱 个"
        evaluated = evaluate(draft)[0]
        evaluated[0]["number"] = 1
        path = self.root / "space.xlsx"
        _write_workbook(path, self.profile, evaluated)
        cell = load_workbook(path, rich_text=True).active["B4"]
        self.assertEqual(cell.value, "收纳箱 个")
        self.assertEqual(cell.font.name, "黑体")

    def test_repeated_names_are_numeric_and_match_excel_and_files(self) -> None:
        sources = [
            ("ride_a.pdf", "10345678901234567890", "客运服务", "交通费", "滴滴网约车1"),
            ("ride_b.pdf", "11345678901234567890", "客运服务", "交通费", "滴滴网约车A"),
            ("ride_c.pdf", "12345678901234567891", "客运服务", "交通费", "滴滴网约车"),
            ("box_a.pdf", "13345678901234567890", "收纳箱", "日用及两用物品", "收纳箱 个"),
            ("box_b.pdf", "14345678901234567890", "整理箱", "日用及两用物品", "整理箱 X-6069 个"),
        ]
        for index, (ref, number, product, _, _) in enumerate(sources, 1):
            (self.input / ref).write_bytes(invoice_pdf(number, f"{index * 10}.00", product))
        draft = self.draft()
        by_ref = {r["ref"]: r for r in draft["records"]}
        for ref, _, _, category, description in sources:
            by_ref[ref]["manual"].update(category=category, description=description)
        package = build(draft)
        state = json.loads((self.output / "state.json").read_text(encoding="utf-8"))
        named = [r for r in state["records"] if r["status"] in {"ready", "missing"}]
        self.assertEqual([r["description"] for r in named if r["category"] == "交通费"],
                         ["滴滴网约车1", "滴滴网约车2", "滴滴网约车3"])
        self.assertEqual([r["description"] for r in named if r["category"] == "日用及两用物品"],
                         ["收纳箱1", "收纳箱2"])
        ws = load_workbook(Path(package["directory"]) / "测试学生-报销.xlsx").active
        excel_names = {str(row[1].value) for row in ws if isinstance(row[0].value, int)}
        self.assertEqual(excel_names, {r["description"] for r in named})
        self.assertTrue(all(f"-{r['description']}-" in r["output_file"] for r in named))

    def test_network_service_needs_verified_subject(self) -> None:
        (self.input / "service.pdf").write_bytes(invoice_pdf("15345678901234567890", "68.00", "网络服务"))
        draft = self.draft()
        draft["records"][0]["manual"].update(category="网络服务费", description="网络服务费")
        self.assertEqual(evaluate(draft)[0][0]["status"], "pending")
        package = build(draft)
        self.assertEqual(package["formal_count"], 1)
        ws = load_workbook(Path(package["directory"]) / "测试学生-报销.xlsx").active
        self.assertEqual(ws["B4"].value, "平台待核实网络服务费")
        self.assertEqual(ws["C4"].value, 68)
        self.assertIn("待核实：网络服务平台", str(ws["D4"].value))
        self.assertTrue(any("平台待核实网络服务费-68元" in p.name for p in (Path(package["directory"]) / "发票").iterdir()))
        self.assertFalse((Path(package["directory"]) / "待核实").exists())
        reminder = "".join(page.extract_text() or "" for page in PdfReader(Path(package["directory"]) / "补件与复核提醒.pdf").pages)
        self.assertIn("Excel", reminder)
        draft["records"][0]["manual"]["service_subject"] = "向日葵"
        record = evaluate(draft)[0][0]
        self.assertEqual(record["status"], "ready")
        self.assertEqual(record["description"], "向日葵网络服务费")


if __name__ == "__main__":
    unittest.main()
